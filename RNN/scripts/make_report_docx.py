import sys, re
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
md=(ROOT/'reports'/'Chuong_4_RNN.md').read_text(encoding='utf-8')
out=ROOT/'reports'/'Chuong_4_RNN_10_15_trang.docx'

def set_cell_shading(cell, fill='D9EAF7'):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def add_table(doc, rows):
    if not rows: return
    data=[]
    for r in rows:
        cells=[c.strip() for c in r.strip().strip('|').split('|')]
        data.append(cells)
    if len(data)>=2 and all(set(c)<=set('-: ') for c in data[1]):
        data.pop(1)
    cols=max(len(r) for r in data)
    table=doc.add_table(rows=len(data), cols=cols)
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.style='Table Grid'
    for i,row in enumerate(data):
        for j in range(cols):
            cell=table.cell(i,j); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
            run=p.add_run(row[j] if j<len(row) else '')
            run.font.name='Times New Roman'; run.font.size=Pt(10)
            if i==0: run.bold=True; set_cell_shading(cell)
    doc.add_paragraph()


doc=Document()
sec=doc.sections[0]
sec.top_margin=Cm(2.0); sec.bottom_margin=Cm(2.0); sec.left_margin=Cm(3.0); sec.right_margin=Cm(2.0)
styles=doc.styles
normal=styles['Normal']; normal.font.name='Times New Roman'; normal.font.size=Pt(13)
normal.paragraph_format.line_spacing=1.5; normal.paragraph_format.space_after=Pt(3); normal.paragraph_format.first_line_indent=Cm(1.0)
for name,size,bold in [('Title',18,True),('Heading 1',15,True),('Heading 2',14,True),('Heading 3',13,True)]:
    st=styles[name]; st.font.name='Times New Roman'; st.font.size=Pt(size); st.font.bold=bold
    st.paragraph_format.space_before=Pt(7); st.paragraph_format.space_after=Pt(4)

lines=md.splitlines(); i=0; in_code=False; code=[]
while i<len(lines):
    line=lines[i]
    if line.startswith('```'):
        if not in_code:
            in_code=True; code=[]
        else:
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(.7); p.paragraph_format.right_indent=Cm(.4); p.paragraph_format.space_after=Pt(5)
            r=p.add_run('\n'.join(code)); r.font.name='Consolas'; r.font.size=Pt(9)
            in_code=False
        i+=1; continue
    if in_code:
        code.append(line); i+=1; continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            rows.append(lines[i]); i+=1
        add_table(doc,rows); continue
    if line.startswith('# '):
        p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run(line[2:].strip()); r.bold=True
    elif line.startswith('## '):
        doc.add_paragraph(line[3:].strip(),style='Heading 1')
    elif line.startswith('### '):
        doc.add_paragraph(line[4:].strip(),style='Heading 2')
    elif line.startswith('#### '):
        doc.add_paragraph(line[5:].strip(),style='Heading 3')
    elif not line.strip():
        pass
    else:
        p=doc.add_paragraph()
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.keep_together=False
        # simple bold markers
        line=line.replace('`','')
        parts=re.split(r'(\*\*.*?\*\*)', line)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                rr=p.add_run(part[2:-2]); rr.bold=True
            else:
                rr=p.add_run(part)
            rr.font.name='Times New Roman'; rr.font.size=Pt(13)
    i+=1

# Footer page number
for section in doc.sections:
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=footer.add_run('Trang '); run.font.name='Times New Roman'; run.font.size=Pt(10)
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); run._r.addnext(fld)

doc.save(out)
print(out)
