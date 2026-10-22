"""Render Chương 3 markdown into a conventional university report DOCX."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
source=ROOT/'reports'/'Chuong_3_CNN.md'
output=ROOT/'reports'/'Chuong_3_CNN_10_15_trang.docx'


def shade(cell,fill):
    tcpr=cell._tc.get_or_add_tcPr();s=OxmlElement('w:shd');s.set(qn('w:fill'),fill);tcpr.append(s)


def add_page_field(par):
    par.add_run('Trang ')
    run=par.add_run()
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
    run._r.addnext(field)


def make():
    d=Document(); sec=d.sections[0]
    sec.page_height=Cm(29.7);sec.page_width=Cm(21)
    sec.top_margin=Cm(2.15);sec.bottom_margin=Cm(2.05);sec.left_margin=Cm(2.6);sec.right_margin=Cm(2.1)
    sec.header_distance=Cm(.8);sec.footer_distance=Cm(.9)
    base=d.styles['Normal'];base.font.name='Times New Roman';base.font.size=Pt(12.1)
    base.paragraph_format.first_line_indent=Cm(.58)
    base.paragraph_format.line_spacing=1.34
    base.paragraph_format.space_after=Pt(5)
    base.paragraph_format.widow_control=True
    for name,size,color,space in [('Heading 1',15,(18,45,71),10),('Heading 2',13,(23,53,80),8),('Heading 3',12.4,(38,67,99),6)]:
        st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor(*color)
        st.paragraph_format.space_before=Pt(space);st.paragraph_format.space_after=Pt(4)
        st.paragraph_format.keep_with_next=True;st.paragraph_format.first_line_indent=Cm(0)
        st.paragraph_format.line_spacing=1.18
    title_style=d.styles.add_style('ChapterTitle', WD_STYLE_TYPE.PARAGRAPH)
    title_style.font.name='Times New Roman';title_style.font.bold=True;title_style.font.size=Pt(16);title_style.font.color.rgb=RGBColor(15,38,70)
    title_style.paragraph_format.space_after=Pt(17);title_style.paragraph_format.line_spacing=1.1
    title_style.paragraph_format.first_line_indent=Cm(0)
    tabcap=d.styles.add_style('TableCaption',WD_STYLE_TYPE.PARAGRAPH)
    tabcap.font.name='Times New Roman';tabcap.font.size=Pt(10.8);tabcap.font.italic=True
    tabcap.paragraph_format.space_before=Pt(5);tabcap.paragraph_format.space_after=Pt(5);tabcap.paragraph_format.first_line_indent=Cm(0)
    hdr=sec.header.paragraphs[0];hdr.text='BÁO CÁO TỔNG QUAN TRÍ TUỆ NHÂN TẠO  |  CHƯƠNG 3'
    hdr.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    hdr.runs[0].font.name='Times New Roman';hdr.runs[0].font.size=Pt(8.5);hdr.runs[0].font.color.rgb=RGBColor(111,125,139)
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(foot)
    for r in foot.runs:r.font.name='Times New Roman';r.font.size=Pt(9)

    def inline_add(par,text):
        # preserve bold segments while keeping human-edited paragraph structure.
        for i,part in enumerate(re.split(r'(\*\*[^*]+\*\*)',text)):
            if not part:continue
            if part.startswith('**') and part.endswith('**'):
                par.add_run(part[2:-2]).bold=True
            else:par.add_run(part)

    lines=source.read_text(encoding='utf-8').splitlines()
    i=0
    while i<len(lines):
        l=lines[i].strip()
        if not l:i+=1;continue
        if l.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                line=lines[i].strip()
                cols=[c.strip() for c in line.strip('|').split('|')]
                if not all(re.fullmatch(r':?-{3,}:?',c or '') for c in cols):rows.append(cols)
                i+=1
            if not rows:continue
            n=len(rows[0]);t=d.add_table(rows=0,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER
            t.style='Table Grid';t.autofit=True
            for ri,row in enumerate(rows):
                cells=t.add_row().cells
                for ci in range(n):
                    text=row[ci] if ci<len(row) else ''
                    cells[ci].text=text.replace('**','')
                    cells[ci].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    for p in cells[ci].paragraphs:
                        p.alignment=WD_ALIGN_PARAGRAPH.CENTER if n>4 else (WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER)
                        p.paragraph_format.space_after=Pt(2)
                        p.paragraph_format.space_before=Pt(2)
                        p.paragraph_format.first_line_indent=Cm(0)
                        p.paragraph_format.line_spacing=1.0
                        if ri < len(rows)-1:
                            p.paragraph_format.keep_with_next = True
                        for run in p.runs:
                            run.font.name='Times New Roman';run.font.size=Pt(8 if n>7 else (10 if n>4 else 10.5));run.font.bold=(ri==0)
                    if ri==0:shade(cells[ci],'DFEAF3')
                    elif ri%2==0:shade(cells[ci],'F7FAFC')
            d.add_paragraph('').paragraph_format.space_after=Pt(1)
            continue
        if l.startswith('# '):
            p=d.add_paragraph(style='ChapterTitle');p.text=l[2:]
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif l.startswith('## '):
            p=d.add_paragraph(l[3:],style='Heading 1')
        elif l.startswith('### '):
            p=d.add_paragraph(l[4:],style='Heading 2')
        elif re.match(r'^Bảng 3\.\d+',l):
            p=d.add_paragraph(l,style='TableCaption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
        elif l.startswith('- '):
            p=d.add_paragraph(style='List Bullet');inline_add(p,l[2:]);p.paragraph_format.first_line_indent=Cm(0)
        else:
            p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;inline_add(p,l)
            if 'https://' in l or re.match(r'^\[\d+\]',l):
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.first_line_indent=Cm(0)
        i+=1
    d.core_properties.title='Chương 3 – Mạng nơ-ron tích chập (CNN)'
    d.core_properties.subject='Báo cáo thực nghiệm CNN From Scratch, Keras, PyTorch trên 3 dataset Kaggle'
    d.save(output)
    print(output,output.stat().st_size)

if __name__=='__main__':make()
